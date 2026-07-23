from __future__ import annotations

import base64
import json
import sys
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

matplotlib.use("Agg")
import matplotlib.pyplot as plt

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from sklearn.ensemble import RandomForestRegressor
from sklearn.linear_model import LinearRegression
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

from src.data_prep import (
    build_modeling_dataset,
    clean_ibovespa_data,
    load_raw_data,
    temporal_train_test_split,
)
from src.forecasting import (
    adf_summary,
    arima_walk_forward_forecast,
    build_regression_dataset,
    evaluate_predictions,
    forecast_next_15_business_days,
    hit_rate,
    run_regression_experiments,
    temporal_split,
)
from src.modeling import FEATURE_COLUMNS


OUTPUT_PATH = PROJECT_ROOT / "reports" / "apresentacao" / "tech_challenge_fase_2_storytelling.docx"
NOTEBOOK_PATH = PROJECT_ROOT / "notebooks" / "tech_challenge_fase_2_ibovespa.ipynb"
FIGURES_DIR = PROJECT_ROOT / "reports" / "figures"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1A1A1A"
MUTED = "666666"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GREEN = "1F7A4D"
GOLD = "7A5A00"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def style_run(run, *, bold=False, size=11, color=INK, font="Calibri") -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.25

    heading_1 = doc.styles["Heading 1"]
    heading_1.font.name = "Calibri"
    heading_1._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    heading_1._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    heading_1.font.size = Pt(16)
    heading_1.font.bold = True
    heading_1.font.color.rgb = RGBColor.from_string(BLUE)
    heading_1.paragraph_format.space_before = Pt(20)
    heading_1.paragraph_format.space_after = Pt(10)

    heading_2 = doc.styles["Heading 2"]
    heading_2.font.name = "Calibri"
    heading_2._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    heading_2._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    heading_2.font.size = Pt(13)
    heading_2.font.bold = True
    heading_2.font.color.rgb = RGBColor.from_string(BLUE)
    heading_2.paragraph_format.space_before = Pt(15)
    heading_2.paragraph_format.space_after = Pt(7.5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Pagina ")
    style_run(run, size=9, color=MUTED)
    add_page_number(footer)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)
    style_run(run, size=9, color=MUTED)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(20 if level == 1 else 15)
    paragraph.paragraph_format.space_after = Pt(10 if level == 1 else 7.5)
    for run in paragraph.runs:
        style_run(
            run,
            bold=True,
            size=16 if level == 1 else 13 if level == 2 else 12,
            color=BLUE if level <= 2 else DARK_BLUE,
        )


def add_paragraph(doc: Document, text: str = ""):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.25
    if text:
        run = paragraph.add_run(text)
        style_run(run)
    return paragraph


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.1666666666666667
    run = paragraph.add_run(text)
    style_run(run)


def add_numbered_item(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.1666666666666667
    run = paragraph.add_run(text)
    style_run(run)


def add_callout(doc: Document, title: str, body: str, color: str = DARK_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p_title = cell.paragraphs[0]
    p_title.paragraph_format.space_after = Pt(3)
    run = p_title.add_run(title)
    style_run(run, bold=True, size=11, color=color)
    p_body = cell.add_paragraph()
    p_body.paragraph_format.space_after = Pt(0)
    run = p_body.add_run(body)
    style_run(run, size=10)
    doc.add_paragraph()


def add_code_block(doc: Document, code: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(code.strip().splitlines()):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        style_run(run, size=9, color="333333", font="Consolas")


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    spacing_after: bool = True,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_row_cant_split(table.rows[0])

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, LIGHT_GRAY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        style_run(run, bold=True, size=9, color=DARK_BLUE)

    for values in rows:
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        for idx, value in enumerate(values):
            cells[idx].width = Inches(widths[idx])
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[idx])
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) <= 18 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            style_run(run, size=9)
    if spacing_after:
        doc.add_paragraph()


def add_picture(doc: Document, path: Path, caption: str, width: float = 6.2) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_p.add_run(caption)
    style_run(caption_run, size=9, color=MUTED)


def fmt_pct(value: float, digits: int = 2) -> str:
    return f"{value * 100:.{digits}f}%".replace(".", ",")


def fmt_int(value: float) -> str:
    return f"{value:,.0f}".replace(",", ".")


def get_notebook() -> dict:
    return json.loads(NOTEBOOK_PATH.read_text(encoding="utf-8"))


def extract_images(nb: dict) -> dict[str, Path]:
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    labels = {
        9: "01_fechamento_historico.png",
        11: "02_volatilidade_retorno.png",
        15: "03_decomposicao.png",
        19: "04_acf_pacf.png",
        25: "05_real_vs_previsto.png",
        36: "06_previsao_15_pregoes.png",
    }
    images: dict[str, Path] = {}
    for cell_index, filename in labels.items():
        for output in nb["cells"][cell_index - 1].get("outputs", []):
            png = output.get("data", {}).get("image/png")
            if png:
                path = FIGURES_DIR / filename
                path.write_bytes(base64.b64decode(png))
                images[filename] = path
                break
    return images


def create_diagnostic_figures(
    test: pd.DataFrame,
    y_true: pd.Series,
    walk_forward_pred: np.ndarray,
) -> dict[str, Path]:
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    actual = np.asarray(y_true, dtype=float)
    predicted = np.asarray(walk_forward_pred, dtype=float)
    error_pct = np.abs(actual - predicted) / actual

    figures: dict[str, Path] = {}

    error_path = FIGURES_DIR / "07_erro_percentual_arima.png"
    fig, ax = plt.subplots(figsize=(10, 5))
    ax.hist(error_pct * 100, bins=18, color="#2E74B5", edgecolor="white")
    ax.axvline(2.15, color="#7A5A00", linestyle="--", linewidth=2, label="Faixa de 2,15%")
    ax.set_title("Distribuicao do erro percentual - ARIMA walk-forward")
    ax.set_xlabel("Erro absoluto percentual (%)")
    ax.set_ylabel("Quantidade de pregoes")
    ax.legend()
    fig.tight_layout()
    fig.savefig(error_path, dpi=150)
    plt.close(fig)
    figures[error_path.name] = error_path

    curve_path = FIGURES_DIR / "08_curva_tolerancia_arima.png"
    tolerances = np.linspace(0.005, 0.035, 13)
    hit_rates = [hit_rate(actual, predicted, tolerance=tolerance) * 100 for tolerance in tolerances]
    fig, ax = plt.subplots(figsize=(10, 5))
    ax.plot(tolerances * 100, hit_rates, marker="o", color="#1F7A4D")
    ax.axvline(2.15, color="#7A5A00", linestyle="--", linewidth=2)
    ax.axhline(80, color="#9B1C1C", linestyle=":", linewidth=2)
    ax.set_title("Acerto por tolerancia de erro - ARIMA walk-forward")
    ax.set_xlabel("Tolerancia de erro (%)")
    ax.set_ylabel("Pregoes dentro da tolerancia (%)")
    ax.set_ylim(0, 105)
    ax.grid(alpha=0.25)
    fig.tight_layout()
    fig.savefig(curve_path, dpi=150)
    plt.close(fig)
    figures[curve_path.name] = curve_path

    error_time_path = FIGURES_DIR / "09_erro_no_tempo_arima.png"
    fig, ax = plt.subplots(figsize=(12, 5))
    ax.plot(test["data"], error_pct * 100, color="#2E74B5", linewidth=1.8)
    ax.axhline(2.15, color="#7A5A00", linestyle="--", linewidth=2)
    ax.set_title("Erro percentual ao longo do teste - ARIMA walk-forward")
    ax.set_xlabel("")
    ax.set_ylabel("Erro absoluto (%)")
    fig.autofmt_xdate()
    fig.tight_layout()
    fig.savefig(error_time_path, dpi=150)
    plt.close(fig)
    figures[error_time_path.name] = error_time_path

    return figures


def run_feature_regression_experiments(model_df: pd.DataFrame):
    train_feat, test_feat = temporal_train_test_split(model_df, test_size=125)
    x_train = train_feat[FEATURE_COLUMNS]
    y_train = train_feat["fechamento_amanha"]
    x_test = test_feat[FEATURE_COLUMNS]
    y_test = test_feat["fechamento_amanha"]

    linear_model = Pipeline(
        [
            ("scaler", StandardScaler()),
            ("model", LinearRegression()),
        ]
    )
    linear_model.fit(x_train, y_train)
    linear_result = evaluate_predictions(
        "Linear Regression - features",
        y_test,
        linear_model.predict(x_test),
    )

    forest_model = RandomForestRegressor(
        n_estimators=300,
        max_depth=6,
        min_samples_leaf=10,
        random_state=42,
        n_jobs=-1,
    )
    forest_model.fit(x_train, y_train)
    forest_result = evaluate_predictions(
        "Random Forest - features",
        y_test,
        forest_model.predict(x_test),
    )
    return [linear_result, forest_result]


def create_model_comparison_figure(results, feature_results) -> Path:
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    path = FIGURES_DIR / "10_comparativo_modelos_regressao.png"
    all_results = list(results) + list(feature_results)
    labels = [item.name.replace("ARIMA(5,1,0) ", "ARIMA ") for item in all_results]
    values = [item.hit_rate_0215 * 100 for item in all_results]
    colors = ["#1F7A4D" if "walk-forward" in label else "#2E74B5" for label in labels]

    fig, ax = plt.subplots(figsize=(11, 5.5))
    ax.barh(labels, values, color=colors)
    ax.axvline(80, color="#9B1C1C", linestyle=":", linewidth=2, label="Meta 80%")
    ax.set_title("Comparativo de acerto dentro da tolerancia de 2,15%")
    ax.set_xlabel("Pregoes dentro da tolerancia (%)")
    ax.set_xlim(0, 105)
    ax.legend(loc="lower right")
    for index, value in enumerate(values):
        ax.text(min(value + 1, 101), index, f"{value:.1f}%", va="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


def build_document() -> Path:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    nb = get_notebook()
    images = extract_images(nb)

    raw = load_raw_data()
    dados = clean_ibovespa_data(raw)
    regression_df = build_regression_dataset()
    model_df = build_modeling_dataset()
    train, test = temporal_split(regression_df, test_size=125)
    results = run_regression_experiments(test_size=125)
    feature_results = run_feature_regression_experiments(model_df)
    adf = adf_summary()
    forecast = forecast_next_15_business_days()
    y_true = test["fechamento_amanha"]
    walk_forward_pred = arima_walk_forward_forecast(train, test)
    images.update(create_diagnostic_figures(test, y_true, walk_forward_pred))
    images["10_comparativo_modelos_regressao.png"] = create_model_comparison_figure(
        results,
        feature_results,
    )

    doc = Document()
    setup_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(100)
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("Previsao do Fechamento do IBOVESPA")
    style_run(run, bold=True, size=28, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(5)
    run = subtitle.add_run("Tech Challenge - Fase 2 - POSTECH Data Analytics")
    style_run(run, size=14, color=INK)
    theme = doc.add_paragraph()
    theme.alignment = WD_ALIGN_PARAGRAPH.CENTER
    theme.paragraph_format.space_after = Pt(150)
    run = theme.add_run("Estrategia de dados, storytelling e modelo preditivo de series temporais")
    style_run(run, size=12, color=INK)
    intro = add_paragraph(
        doc,
        "Este documento descreve, passo a passo, a estrategia de dados construida para prever o fechamento diario do indice IBOVESPA, incluindo a explicacao de cada etapa do codigo desenvolvido, as decisoes tecnicas tomadas e os resultados obtidos. O codigo completo e executavel esta disponivel no notebook Jupyter que acompanha esta entrega.",
    )
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    meta_name = doc.add_paragraph()
    meta_name.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    meta_name.paragraph_format.space_after = Pt(10)
    meta_name.paragraph_format.line_spacing = 1.25
    run = meta_name.add_run("Nome: Renan Cotrin Lapastina")
    style_run(run, bold=True, size=11, color=INK)
    meta_rm = doc.add_paragraph()
    meta_rm.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    meta_rm.paragraph_format.space_after = Pt(10)
    meta_rm.paragraph_format.line_spacing = 1.25
    run = meta_rm.add_run("RM: 366292")
    style_run(run, bold=True, size=11, color=INK)
    add_callout(
        doc,
        "Resumo executivo",
        "A entrega trata o problema como serie temporal e regressao. A assertividade principal e calculada como 100 * (1 - WMAPE), conforme orientacao recebida para uso de metricas de distancia.",
        color=GREEN,
    )
    doc.add_page_break()

    add_heading(doc, "1. Escopo analitico e criterio de sucesso")
    add_paragraph(
        doc,
        "O desafio foi construir um modelo preditivo de series temporais para o fechamento diario do IBOVESPA, com assertividade minima de 80% e uso como ferramenta academica de apoio a decisao. Antes de escolher modelos, foi necessario definir o que significa acertar uma previsao em uma serie financeira.",
    )
    add_paragraph(
        doc,
        "O fechamento de um indice amplo costuma ficar muito proximo do fechamento anterior. Por isso, uma previsao numerica pode apresentar erro percentual baixo mesmo quando o mercado continua dificil de antecipar. A estrategia adotada separa duas perguntas: prever o valor de fechamento, que e a entrega formal por regressao, e investigar a direcao do proximo pregao como analise complementar.",
    )
    add_bullet(
        doc,
        "Regressao: prever o fechamento do proximo pregao e avaliar com WMAPE, MAE, assertividade por 1 - WMAPE e acerto dentro de uma faixa de tolerancia.",
    )
    add_bullet(
        doc,
        "Classificacao: avaliar alta ou baixa como investigacao paralela, sem forcar uma acuracia artificial em uma tarefa naturalmente mais instavel.",
    )
    add_callout(
        doc,
        "Decisao metodologica",
        "A meta formal e defendida pela regressao, com validacao temporal correta. A classificacao entra como transparencia tecnica sobre os limites reais do problema.",
        color=GREEN,
    )
    add_heading(doc, "1.1 Aderencia ao enunciado", level=2)
    add_table(
        doc,
        ["Exigencia da prova", "Como foi atendida"],
        [
            ["Modelo de serie temporal", "ARIMA(5,1,0) com validacao walk-forward e comparacao contra baseline."],
            ["Assertividade minima de 80%", "98,67% por 1 - WMAPE e 82,40% dentro da faixa de 2,15%."],
            ["Storytelling dos picos", "Linha do tempo economica e leitura dos choques de volatilidade."],
            ["Explicacao e vantagens do modelo", "ADF, ACF/PACF, vantagens do ARIMA e comparacao com modelos alternativos."],
            ["Previsao de 15 dias", "Tabela e grafico de 23/07/2026 a 12/08/2026 com intervalo de confianca."],
        ],
        [2.3, 4.2],
        spacing_after=False,
    )

    doc.add_page_break()
    add_heading(doc, "1.2 Como a entrega esta organizada", level=2)
    for item in [
        "Leitura, limpeza e padronizacao da base historica.",
        "Storytelling exploratorio e contexto economico dos ciclos do IBOVESPA.",
        "Decomposicao da serie, teste ADF e leitura de ACF/PACF.",
        "Engenharia de atributos para modelos complementares.",
        "Separacao treino/teste respeitando a ordem temporal.",
        "Baseline de persistencia, ARIMA estatico e ARIMA walk-forward.",
        "Comparacao de resultados por WMAPE, MAE e tolerancia de erro.",
        "Predicao dos proximos 15 pregoes e conclusao com limitacoes.",
    ]:
        add_bullet(doc, item)
    doc.add_page_break()

    add_heading(doc, "2. Preparacao da base historica")
    add_paragraph(
        doc,
        "A base utilizada contem o historico diario do IBOVESPA obtido no Investing.com, com data, fechamento, abertura, maxima, minima, volume negociado e variacao percentual. O arquivo original vem em ordem decrescente, do pregao mais recente para o mais antigo, e usa formato brasileiro nos numeros.",
    )
    add_paragraph(
        doc,
        "A etapa de limpeza foi tratada com cuidado porque valores como 175.831 representam 175.831 pontos, nao 175,831 pontos. Por isso, a leitura inicial preserva tudo como texto e a conversao numerica acontece de forma controlada.",
    )
    add_table(
        doc,
        ["Indicador", "Valor"],
        [
            ["Periodo bruto", f"{dados['data'].min().date()} a {dados['data'].max().date()}"],
            ["Quantidade de pregoes", fmt_int(len(dados))],
            ["Colunas originais", f"{raw.shape[1]}"],
            ["Fonte", "Investing.com - historico diario do IBOVESPA"],
        ],
        [2.2, 4.3],
    )
    add_heading(doc, "2.1 Regras de limpeza implementadas", level=2)
    for item in [
        "Conversao da coluna de data para tipo datetime e reordenacao cronologica.",
        "Conversao de fechamento, abertura, maxima e minima do formato brasileiro para numero.",
        "Conversao da variacao percentual removendo o simbolo de porcentagem e respeitando virgula decimal.",
        "Padronizacao do volume negociado, que mistura unidades K, M e B no export do Investing.com.",
        "Normalizacao dos nomes de colunas para evitar problemas com acentos em diferentes ambientes.",
    ]:
        add_bullet(doc, item)
    add_code_block(
        doc,
        """
raw_df = load_raw_data()
dados = clean_ibovespa_data(raw_df)

cleaned["data"] = pd.to_datetime(cleaned["data"], format="%d.%m.%Y")
cleaned = cleaned.sort_values("data").reset_index(drop=True)

for column in ["fechamento", "abertura", "maxima", "minima"]:
    cleaned[column] = cleaned[column].apply(parse_brazilian_number)
""",
    )
    add_paragraph(
        doc,
        "O volume recebeu uma funcao propria porque a fonte alterna unidades no mesmo campo. Sem esse cuidado, valores em milhoes e bilhoes poderiam ser comparados como se estivessem na mesma escala, distorcendo qualquer atributo derivado de liquidez.",
    )

    add_heading(doc, "3. Leitura historica do IBOVESPA")
    add_paragraph(
        doc,
        "Com os dados limpos, a primeira leitura foi visual: antes de propor modelos, e preciso entender o comportamento historico do indice. A serie do IBOVESPA nao apresenta um padrao sazonal simples; ela e marcada por ciclos longos, choques macroeconomicos e mudancas de regime.",
    )
    if "01_fechamento_historico.png" in images:
        add_picture(doc, images["01_fechamento_historico.png"], "Figura 1 - Evolucao historica do fechamento do IBOVESPA.")
    add_heading(doc, "3.1 Ciclos economicos observados", level=2)
    add_paragraph(
        doc,
        "A leitura de negocio foi organizada em periodos. Essa etapa e importante porque transforma o grafico em storytelling: cada movimento relevante da serie passa a ser conectado a um contexto economico e politico plausivel.",
    )
    add_table(
        doc,
        ["Periodo", "Comportamento", "Contexto"],
        [
            ["2010-2015", "Queda / estagnacao", "Desaceleracao economica, deterioracao fiscal e perda de confianca do mercado."],
            ["2015-2016", "Ponto de minimo", "Crise politica, impeachment e recessao economica elevaram a aversao a risco."],
            ["2016-2020", "Ciclo de alta", "Expectativa de reformas, juros globais baixos e retomada gradual do fluxo para risco."],
            ["Fev-Mar/2020", "Choque abrupto", "Pandemia de Covid-19, circuit breakers e repricing global de ativos."],
            ["2020-2021", "Recuperacao em V", "Estimulos monetarios/fiscais e Selic em patamar historicamente baixo."],
            ["2021-2022", "Lateralizacao / queda", "Aperto monetario para conter inflacao e aumento da incerteza eleitoral."],
            ["2023-2026", "Novo ciclo de alta", "Ciclo de corte de juros, resultados corporativos e renovacao de recordes do indice."],
        ],
        [1.3, 1.7, 3.5],
    )
    add_paragraph(
        doc,
        "Esse pano de fundo reforca que o IBOVESPA precisa ser reavaliado periodicamente: o comportamento estatistico do indice muda quando o ambiente de juros, risco politico, fluxo estrangeiro e resultados corporativos tambem muda.",
    )
    add_heading(doc, "3.2 Retornos, choques e faixa de tolerancia", level=2)
    add_paragraph(
        doc,
        "A variacao diaria mostra que a maior parte dos pregoes se concentra em movimentos pequenos, enquanto choques extremos aparecem em momentos especificos, especialmente em 2020. Essa distribuicao justifica usar uma metrica de tolerancia alem do WMAPE.",
    )
    if "02_volatilidade_retorno.png" in images:
        add_picture(doc, images["02_volatilidade_retorno.png"], "Figura 2 - Retornos diarios e distribuicao da variacao percentual.")
    add_callout(
        doc,
        "Faixa operacional",
        "A media de variacao absoluta diaria ficou proxima de 1,04%, enquanto o percentil 90 ficou perto de 2,20%. Por isso, a faixa de 2,15% foi usada como criterio pratico de acerto alem do WMAPE.",
        color=GOLD,
    )
    doc.add_page_break()

    add_heading(doc, "4. Diagnostico estatistico da serie")
    add_heading(doc, "4.1 Separacao entre tendencia, sazonalidade e residuo", level=2)
    add_paragraph(
        doc,
        "Para decompor a serie em tendencia, sazonalidade e residuo, a serie foi reorganizada em frequencia diaria e preenchida com o ultimo fechamento conhecido apenas para fins visuais. Os modelos preditivos continuam usando os dias reais de pregao.",
    )
    if "03_decomposicao.png" in images:
        add_picture(doc, images["03_decomposicao.png"], "Figura 3 - Decomposicao da serie de fechamento.")
    add_paragraph(
        doc,
        "A decomposicao mostra que a tendencia explica a maior parte do movimento historico. A sazonalidade anual aparece com amplitude pequena, enquanto os residuos concentram choques nao explicados, principalmente em momentos de estresse de mercado.",
    )
    doc.add_page_break()
    add_heading(doc, "4.2 Estacionariedade pelo teste ADF", level=2)
    add_paragraph(
        doc,
        "O teste ADF foi usado para decidir a ordem de diferenciacao do ARIMA. A hipotese nula do teste e que a serie possui raiz unitaria, ou seja, nao e estacionaria.",
    )
    add_table(
        doc,
        ["Teste ADF", "P-valor", "Interpretacao"],
        [
            ["Serie original", f"{adf['p_valor_original']:.4f}".replace(".", ","), "Nao estacionaria"],
            ["Primeira diferenca", f"{adf['p_valor_primeira_diferenca']:.8f}".replace(".", ","), "Estacionaria"],
        ],
        [2.0, 1.6, 2.9],
    )
    add_paragraph(
        doc,
        "Como a serie original nao e estacionaria e a primeira diferenca se torna estacionaria, a especificacao integrada d=1 e tecnicamente justificada.",
    )
    add_heading(doc, "4.3 Memoria curta pela ACF e PACF", level=2)
    if "04_acf_pacf.png" in images:
        add_picture(
            doc,
            images["04_acf_pacf.png"],
            "Figura 4 - ACF e PACF usadas como apoio para escolha do ARIMA.",
            width=5.55,
        )
    add_paragraph(
        doc,
        "Os graficos de autocorrelacao indicam dependencia de curto prazo, sugerindo que os ultimos pregoes carregam informacao relevante para a previsao do proximo fechamento. Essa leitura sustenta o teste de um ARIMA com componente autorregressivo curto.",
    )
    doc.add_page_break()

    add_heading(doc, "5. Construcao de alvos e atributos")
    add_paragraph(
        doc,
        "Embora o ARIMA seja a entrega principal, tambem foram criados atributos tecnicos para modelos supervisionados de comparacao. Todos os atributos usam apenas informacao disponivel ate o dia da previsao, evitando vazamento de dados.",
    )
    for item in [
        "Retornos acumulados e defasados em janelas curtas.",
        "Medias moveis e distancia do fechamento em relacao as medias.",
        "Volatilidade recente, amplitude intradiaria e corpo do candle.",
        "Posicao do fechamento em janelas recentes.",
        "Indicadores tecnicos como RSI, MACD e Bandas de Bollinger.",
        "Variaveis simples de calendario, como dia da semana e mes.",
    ]:
        add_bullet(doc, item)
    add_table(
        doc,
        ["Indicador", "Valor"],
        [
            ["Base com features", f"{fmt_int(len(model_df))} registros"],
            ["Inicio da base modelavel", str(model_df["data"].min().date())],
            ["Fim da base modelavel", str(model_df["data"].max().date())],
            ["Proporcao historica de altas", fmt_pct(float(model_df["alta_amanha"].mean()))],
        ],
        [2.6, 3.9],
    )
    add_code_block(
        doc,
        """
featured["fechamento_amanha"] = featured["fechamento"].shift(-1)
featured["alta_amanha"] = (
    featured["fechamento_amanha"] > featured["fechamento"]
).astype(int)

featured["retorno_5d"] = featured["fechamento"].pct_change(5)
featured["media_movel_21"] = featured["fechamento"].rolling(21).mean()
featured["volatilidade_21d"] = featured["retorno_1d"].rolling(21).std()
""",
    )
    add_paragraph(
        doc,
        "O ponto central dessa etapa e que o alvo sempre olha para o proximo pregao, enquanto as features olham apenas para tras. Essa separacao reduz o risco de lookahead bias e deixa a avaliacao mais parecida com o uso real: ao final do dia D, o modelo tenta estimar D+1.",
    )
    doc.add_page_break()

    add_heading(doc, "6. Validacao fora da amostra")
    add_paragraph(
        doc,
        "Para evitar vazamento de informacao futura, a divisao entre treino e teste foi feita de forma temporal. A ultima observacao da base limpa, 22/07/2026, e usada como fechamento conhecido; como o alvo e o proximo pregao, o ultimo alvo conhecido fica em 21/07/2026.",
    )
    add_table(
        doc,
        ["Conjunto", "Registros", "Periodo"],
        [
            ["Treino", fmt_int(len(train)), f"{train['data'].min().date()} a {train['data'].max().date()}"],
            ["Teste", fmt_int(len(test)), f"{test['data'].min().date()} a {test['data'].max().date()}"],
        ],
        [1.4, 1.3, 3.8],
    )
    add_code_block(
        doc,
        """
def temporal_split(dataset, test_size=125):
    train = dataset.iloc[:-test_size].copy()
    test = dataset.iloc[-test_size:].copy()
    return train, test
""",
    )
    add_callout(
        doc,
        "Validacao temporal",
        "Nao houve embaralhamento aleatorio. O modelo foi treinado no passado e avaliado no periodo mais recente com alvo conhecido, simulando melhor uma situacao real de previsao.",
        color=DARK_BLUE,
    )
    doc.add_page_break()

    add_heading(doc, "7. Experimentos e escolha do modelo")
    add_paragraph(
        doc,
        "Foram comparados tres caminhos: baseline de persistencia, ARIMA estatico e ARIMA walk-forward. A persistencia usa o fechamento de hoje como previsao para amanha; o ARIMA estatico treina uma vez; o walk-forward recalibra o historico a cada novo ponto observado no teste.",
    )
    add_paragraph(
        doc,
        "A funcao de avaliacao concentra as metricas principais. O WMAPE foi escolhido por medir erro percentual ponderado pelo tamanho real da serie; em seguida, a assertividade e calculada por 1 - WMAPE, conforme orientacao recebida para metricas de distancia.",
    )
    add_code_block(
        doc,
        """
def wmape(y_true, y_pred):
    actual = np.asarray(y_true, dtype=float)
    predicted = np.asarray(y_pred, dtype=float)
    return np.sum(np.abs(actual - predicted)) / np.sum(np.abs(actual))

assertiveness = 1 - wmape(y_true, y_pred)
""",
    )
    add_heading(doc, "7.1 Referencia ingenua: fechamento de hoje", level=2)
    add_paragraph(
        doc,
        "O primeiro modelo testado foi o baseline Naive, ou persistencia: a previsao para amanha e simplesmente o fechamento de hoje. Em series financeiras de alta liquidez, esse baseline costuma ser muito competitivo em erro percentual, justamente porque a variacao diaria tende a ser pequena.",
    )
    add_code_block(
        doc,
        """
naive_pred = test_reg["fechamento"].values
naive_result = evaluate_predictions(
    "Naive - persistencia",
    y_true,
    naive_pred,
)
""",
    )
    naive_result = next(item for item in results if item.name == "Naive - persistencia")
    add_table(
        doc,
        ["Metrica do baseline", "Resultado", "Leitura"],
        [
            ["WMAPE", fmt_pct(naive_result.wmape), "Erro percentual ponderado muito baixo."],
            ["Assertividade", fmt_pct(naive_result.assertiveness), "Acima da meta, mas baseada em regra ingenua."],
            ["MAE", f"{fmt_int(naive_result.mae)} pontos", "Erro medio absoluto em pontos do indice."],
            ["Acerto 2,15%", fmt_pct(naive_result.hit_rate_0215), "Referencia operacional para comparar modelos."],
        ],
        [1.7, 1.4, 3.4],
    )
    add_paragraph(
        doc,
        "Esse baseline nao deve ser ignorado. Ele funciona como referencia minima: qualquer modelo mais sofisticado precisa ser comparado contra ele para que o resultado nao pareca melhor do que realmente e. Ao mesmo tempo, ele nao oferece estrutura estatistica, intervalo de confianca nem explicacao de comportamento da serie; ele apenas repete o ultimo fechamento conhecido.",
    )

    add_heading(doc, "7.2 Racional estatistico para o ARIMA", level=2)
    add_paragraph(
        doc,
        "A escolha do ARIMA nao foi feita apenas por tentativa e erro. O teste ADF mostrou que a serie original nao e estacionaria, enquanto a primeira diferenca se torna estacionaria. Isso sustenta o termo integrado d=1. Os graficos de ACF e PACF indicam dependencia de curto prazo, coerente com uma especificacao ARIMA simples e interpretavel.",
    )
    add_table(
        doc,
        ["Vantagem", "Por que importa neste projeto"],
        [
            ["Fundamentacao estatistica", "A ordem do modelo e apoiada por ADF, ACF e PACF, tornando a decisao auditavel."],
            ["Interpretabilidade", "Os componentes autorregressivo e integrado sao explicaveis para uma banca tecnica."],
            ["Intervalo de confianca", "A previsao traz faixa de incerteza, algo essencial em series financeiras."],
            ["Baixo custo computacional", "Funciona bem com uma unica serie historica de cerca de 4 mil pregoes."],
            ["Extensibilidade", "Pode evoluir para SARIMAX com Selic, cambio ou indices externos como variaveis exogenas."],
        ],
        [2.0, 4.5],
    )
    add_heading(doc, "7.3 Validacao walk-forward", level=2)
    add_paragraph(
        doc,
        "Na avaliacao walk-forward, o modelo e recalibrado a cada passo com o fechamento real observado no periodo de teste. Isso aproxima a avaliacao de um uso pratico: no fim de cada pregao, o historico e atualizado e uma nova previsao e gerada.",
    )
    add_code_block(
        doc,
        """
history = list(train["fechamento"].values)
predictions = []
for actual_close in test["fechamento"].values:
    model = ARIMA(history, order=(5, 1, 0)).fit()
    predictions.append(float(model.forecast(steps=1)[0]))
    history.append(actual_close)
""",
    )

    add_heading(doc, "7.4 Resultado comparativo", level=2)
    rows = []
    for result in results:
        rows.append(
            [
                result.name,
                fmt_pct(result.wmape),
                fmt_pct(result.assertiveness),
                f"{fmt_int(result.mae)} pontos",
                fmt_pct(result.hit_rate_0215),
            ]
        )
    add_table(
        doc,
        ["Modelo", "WMAPE", "Assertividade", "MAE", "Acerto 2,15%"],
        rows,
        [2.2, 1.0, 1.1, 1.1, 1.1],
    )
    if "05_real_vs_previsto.png" in images:
        add_picture(doc, images["05_real_vs_previsto.png"], "Figura 5 - Fechamento real versus previsoes no periodo de teste.")
    if "09_erro_no_tempo_arima.png" in images:
        add_picture(doc, images["09_erro_no_tempo_arima.png"], "Figura 6 - Erro percentual do ARIMA walk-forward ao longo do periodo de teste.")
    doc.add_page_break()
    add_callout(
        doc,
        "Modelo escolhido",
        "O ARIMA(5,1,0) walk-forward foi escolhido para a narrativa principal por respeitar a estrutura temporal e atingir 98,67% de assertividade por 1 - WMAPE, alem de 82,40% de acerto dentro da faixa de 2,15%.",
        color=GREEN,
    )
    add_heading(doc, "7.5 Regressao com atributos tecnicos", level=2)
    add_paragraph(
        doc,
        "Como contraponto ao ARIMA, foram testados modelos de regressao supervisionada com os 37 atributos tecnicos construidos na etapa de engenharia. A ideia foi verificar se sinais de retorno, volatilidade, medias moveis, candle, volume e calendario melhorariam a previsao do fechamento seguinte.",
    )
    add_code_block(
        doc,
        """
lin_reg = Pipeline([
    ("scaler", StandardScaler()),
    ("model", LinearRegression()),
])
lin_reg.fit(x_train, y_train)

rf_reg = RandomForestRegressor(
    n_estimators=300,
    max_depth=6,
    min_samples_leaf=10,
    random_state=42,
)
rf_reg.fit(x_train, y_train)
""",
    )
    feature_rows = []
    for result in feature_results:
        feature_rows.append(
            [
                result.name,
                fmt_pct(result.wmape),
                fmt_pct(result.assertiveness),
                f"{fmt_int(result.mae)} pontos",
                fmt_pct(result.hit_rate_0215),
            ]
        )
    add_table(
        doc,
        ["Modelo com atributos", "WMAPE", "Assertividade", "MAE", "Acerto 2,15%"],
        feature_rows,
        [2.1, 1.0, 1.1, 1.1, 1.2],
    )
    add_paragraph(
        doc,
        "A regressao linear fica competitiva porque o proprio fechamento atual esta entre as variaveis de entrada, fazendo o modelo aprender uma relacao muito proxima da persistencia. O Random Forest, por outro lado, perde aderencia neste recorte porque modelos baseados em arvores tendem a criar faixas de previsao e podem ter dificuldade para extrapolar suavemente o nivel atual do indice.",
    )
    add_paragraph(
        doc,
        "Por isso, os modelos com atributos entram como validacao adicional, nao como escolha final. Eles mostram que adicionar muitos indicadores tecnicos nao garante ganho real sobre uma serie financeira de curto prazo. O ARIMA walk-forward permanece como modelo principal por ser mais aderente ao enunciado de series temporais, ter fundamentacao estatistica pela estacionariedade/autocorrelacao e entregar intervalo de confianca para a previsao de 15 pregoes.",
    )
    if "10_comparativo_modelos_regressao.png" in images:
        add_picture(doc, images["10_comparativo_modelos_regressao.png"], "Figura 7 - Comparacao entre ARIMA, baseline e regressoes com atributos pela tolerancia de 2,15%.")

    add_heading(doc, "7.6 Robustez da tolerancia de erro", level=2)
    add_paragraph(
        doc,
        "Como a leitura de acerto por tolerancia depende da margem escolhida, tambem foi feita uma analise de sensibilidade. Isso evita escolher uma unica faixa sem contexto e mostra como o modelo se comporta em margens mais conservadoras e mais flexiveis.",
    )
    tolerance_rows = []
    for tolerance in [0.010, 0.015, 0.020, 0.0215, 0.025, 0.030]:
        tolerance_rows.append(
            [
                fmt_pct(tolerance),
                fmt_pct(hit_rate(y_true, walk_forward_pred, tolerance=tolerance)),
            ]
        )
    add_table(
        doc,
        ["Tolerancia de erro", "Acerto ARIMA walk-forward"],
        tolerance_rows,
        [2.4, 2.4],
    )
    if "07_erro_percentual_arima.png" in images:
        add_picture(doc, images["07_erro_percentual_arima.png"], "Figura 8 - Distribuicao do erro percentual absoluto do ARIMA walk-forward.")
    if "08_curva_tolerancia_arima.png" in images:
        add_picture(doc, images["08_curva_tolerancia_arima.png"], "Figura 9 - Sensibilidade do acerto conforme a tolerancia de erro.")

    doc.add_page_break()

    add_heading(doc, "8. Investigacao complementar de direcao")
    add_paragraph(
        doc,
        "Tambem foram testados modelos supervisionados com features tecnicas e classificacao de direcao. Essa frente foi mantida como complemento, pois a orientacao atual da prova favorece regressao e metricas de distancia. A classificacao de alta/baixa e naturalmente mais instavel, especialmente em dias com variacoes pequenas.",
    )
    add_paragraph(
        doc,
        "Essa decisao deixa a entrega mais honesta. Em vez de inflar uma acuracia de direcao por escolha oportunista de janela, threshold ou hiperparametro, o projeto reporta a regressao como resposta principal e a classificacao como investigacao dos limites do problema.",
    )
    add_callout(
        doc,
        "Leitura tecnica",
        "Prever o valor aproximado do fechamento e diferente de prever corretamente a direcao diaria. A segunda tarefa e mais dura e tende a ficar perto de 50% quando usamos apenas informacoes publicas do proprio indice.",
        color=GOLD,
    )
    doc.add_page_break()

    add_heading(doc, "9. Projecao dos proximos 15 pregoes")
    add_paragraph(
        doc,
        "Com a base atualizada ate 22/07/2026, a previsao inicia em 23/07/2026 e termina em 12/08/2026. O intervalo de confianca aumenta ao longo do horizonte, refletindo a incerteza acumulada.",
    )
    forecast_rows = []
    for _, row in forecast.iterrows():
        forecast_rows.append(
            [
                row["data"].strftime("%d/%m/%Y"),
                fmt_int(row["fechamento_previsto"]),
                fmt_int(row["ic_inferior_95"]),
                fmt_int(row["ic_superior_95"]),
            ]
        )
    add_table(
        doc,
        ["Data", "Previsto", "IC 95% inf.", "IC 95% sup."],
        forecast_rows,
        [1.3, 1.6, 1.8, 1.8],
    )
    if "06_previsao_15_pregoes.png" in images:
        add_picture(doc, images["06_previsao_15_pregoes.png"], "Figura 10 - Projecao dos proximos 15 pregoes.")

    doc.add_page_break()

    add_heading(doc, "10. Fechamento tecnico")
    add_paragraph(
        doc,
        "A solucao final atende ao objetivo da prova substitutiva por tratar o IBOVESPA como serie temporal, usar decomposicao e estacionariedade para justificar o ARIMA, avaliar previsao numerica com WMAPE e apresentar uma previsao de 15 pregoes com intervalo de confianca. O resultado principal supera a meta de 80% pela metrica de assertividade orientada por erro percentual.",
    )
    add_heading(doc, "10.1 Principais entregas", level=2)
    for item in [
        "Uma base historica atualizada e limpa do IBOVESPA ate 22/07/2026.",
        "Storytelling conectando ciclos do indice a eventos economicos relevantes.",
        "Decomposicao, teste ADF e leitura de autocorrelacao para justificar a abordagem temporal.",
        "Comparacao entre baseline, ARIMA estatico e ARIMA walk-forward.",
        "Previsao de 15 pregoes com intervalo de confianca.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "10.2 Limites da abordagem e evolucoes", level=2)
    add_paragraph(
        doc,
        "Uma assertividade alta por WMAPE nao significa prever perfeitamente o mercado. O proprio baseline de persistencia fica forte porque o IBOVESPA geralmente nao muda muitos pontos de um pregao para o outro. Por isso, a entrega combina WMAPE, MAE, tolerancia de erro e interpretacao economica, em vez de vender uma unica metrica como verdade absoluta.",
    )
    add_paragraph(
        doc,
        "A classificacao de direcao diaria segue mais dificil e instavel, o que e coerente com a literatura de mercados eficientes. Como proximo passo, o projeto poderia evoluir para SARIMAX, incorporando variaveis exogenas como Selic, cambio, juros futuros, S&P 500, commodities e indicadores de sentimento.",
    )
    add_paragraph(
        doc,
        "A interpretacao deve permanecer academica e analitica: o modelo nao substitui avaliacao de risco, eventos macroeconomicos ou analise de mercado em tempo real.",
    )

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_document()
    print(path)
